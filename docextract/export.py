"""Write results as JSON, CSV or Excel.

Excel gets three sheets: one row per document, one row per line item, and the review issues,
with low-confidence cells highlighted so a reviewer sees at a glance what to check.
"""

import json
from pathlib import Path

import pandas as pd

from .schema import HEADER_FIELDS, ExtractionResult


def to_json(results: list[ExtractionResult]) -> str:
    return json.dumps([r.model_dump(exclude={"source_text"}) for r in results], indent=2, ensure_ascii=False)


def documents_frame(results: list[ExtractionResult]) -> pd.DataFrame:
    rows = []
    for r in results:
        row = {"source": r.source, "method": r.method, "needs_review": r.needs_review}
        row.update({f: getattr(r.data, f) for f in HEADER_FIELDS})
        row["n_line_items"] = len(r.data.line_items)
        row["issues"] = "; ".join(i.message for i in r.issues)
        row["seconds"] = r.seconds
        rows.append(row)
    return pd.DataFrame(rows)


def line_items_frame(results: list[ExtractionResult]) -> pd.DataFrame:
    rows = [
        {"source": r.source, "method": r.method, "line": i + 1, **li.model_dump()}
        for r in results
        for i, li in enumerate(r.data.line_items)
    ]
    return pd.DataFrame(rows, columns=["source", "method", "line", "description", "quantity", "unit_price", "amount"])


def issues_frame(results: list[ExtractionResult]) -> pd.DataFrame:
    rows = [{"source": r.source, "method": r.method, **i.model_dump()} for r in results for i in r.issues]
    return pd.DataFrame(rows, columns=["source", "method", "field", "code", "severity", "message"])


def to_csv(results: list[ExtractionResult], path: str | Path) -> None:
    documents_frame(results).to_csv(path, index=False)


def to_excel(results: list[ExtractionResult], path: str | Path, review_confidence: float = 0.6) -> None:
    from openpyxl.styles import PatternFill

    docs = documents_frame(results)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        docs.to_excel(xw, sheet_name="Documents", index=False)
        line_items_frame(results).to_excel(xw, sheet_name="Line items", index=False)
        issues_frame(results).to_excel(xw, sheet_name="Issues", index=False)

        # Highlight low-confidence and flagged cells on the Documents sheet.
        ws = xw.sheets["Documents"]
        amber = PatternFill("solid", fgColor="FFE699")
        red = PatternFill("solid", fgColor="F4B6B6")
        cols = {name: idx + 1 for idx, name in enumerate(docs.columns)}
        for row_idx, r in enumerate(results, start=2):
            flagged = {i.field for i in r.issues if i.severity == "error"}
            for f in HEADER_FIELDS:
                cell = ws.cell(row=row_idx, column=cols[f])
                if f in flagged:
                    cell.fill = red
                elif r.confidence.get(f, 1.0) < review_confidence:
                    cell.fill = amber
            if r.needs_review:
                ws.cell(row=row_idx, column=cols["needs_review"]).fill = red
        for ws in xw.sheets.values():  # readable column widths
            for col in ws.columns:
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 60)


def save(results: list[ExtractionResult], path: str | Path) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    ext = path.suffix.lower()
    if ext == ".json":
        path.write_text(to_json(results), encoding="utf-8")
    elif ext == ".csv":
        to_csv(results, path)
    elif ext in (".xlsx", ".xls"):
        to_excel(results, path.with_suffix(".xlsx"))
    else:
        raise ValueError(f"Unsupported export format: {ext} (use .json, .csv or .xlsx)")
